import io
from unittest.mock import AsyncMock, patch

import openpyxl
import pytest
from httpx import AsyncClient

BASE = "/api/v1"
MONTH = "2026-05"


async def _auth(client: AsyncClient, suffix: str) -> dict:
    await client.post(f"{BASE}/auth/register", json={
        "email": f"exp{suffix}@test.com",
        "password": "password123",
        "display_name": f"Export {suffix}",
    })
    resp = await client.post(f"{BASE}/auth/login", json={
        "email": f"exp{suffix}@test.com", "password": "password123",
    })
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


async def _first_category_id(client: AsyncClient, headers: dict) -> str:
    cats = (await client.get(f"{BASE}/categories", headers=headers)).json()
    return cats[0]["id"]


async def _seed_expenses(client: AsyncClient, headers: dict, cat_id: str) -> None:
    for i, (title, amount) in enumerate([("Loyer", "3000"), ("Courses", "450"), ("Internet", "200")]):
        await client.post(f"{BASE}/expenses", json={
            "title": title, "amount": amount, "currency": "MAD",
            "date": f"2026-05-{i + 1:02d}", "category_id": cat_id,
        }, headers=headers)


# ── Excel ─────────────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_excel_export_returns_correct_content_type(client: AsyncClient):
    h = await _auth(client, "xl1")
    resp = await client.get(f"{BASE}/expenses/export", params={"format": "xlsx", "month": MONTH}, headers=h)
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]


@pytest.mark.asyncio
async def test_excel_export_content_disposition(client: AsyncClient):
    h = await _auth(client, "xl2")
    resp = await client.get(f"{BASE}/expenses/export", params={"format": "xlsx", "month": MONTH}, headers=h)
    assert MONTH in resp.headers["content-disposition"]
    assert resp.headers["content-disposition"].endswith(".xlsx\"")


@pytest.mark.asyncio
async def test_excel_export_is_valid_workbook(client: AsyncClient):
    h = await _auth(client, "xl3")
    resp = await client.get(f"{BASE}/expenses/export", params={"format": "xlsx", "month": MONTH}, headers=h)
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.active is not None


@pytest.mark.asyncio
async def test_excel_export_contains_expense_rows(client: AsyncClient):
    h = await _auth(client, "xl4")
    cat_id = await _first_category_id(client, h)
    await _seed_expenses(client, h, cat_id)

    resp = await client.get(f"{BASE}/expenses/export", params={"format": "xlsx", "month": MONTH}, headers=h)
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active

    # Row 1 is header; rows 2+ are data; last row is total footer
    data_rows = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row)]
    assert "Loyer" in data_rows
    assert "Courses" in data_rows
    assert "Internet" in data_rows


@pytest.mark.asyncio
async def test_excel_export_total_row_present(client: AsyncClient):
    h = await _auth(client, "xl5")
    cat_id = await _first_category_id(client, h)
    await _seed_expenses(client, h, cat_id)

    resp = await client.get(f"{BASE}/expenses/export", params={"format": "xlsx", "month": MONTH}, headers=h)
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active

    # Last row col B should be "TOTAL"
    last_row = ws.max_row
    assert ws.cell(row=last_row, column=2).value == "TOTAL"


@pytest.mark.asyncio
async def test_excel_export_empty_month(client: AsyncClient):
    h = await _auth(client, "xl6")
    resp = await client.get(f"{BASE}/expenses/export", params={"format": "xlsx", "month": "2020-01"}, headers=h)
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    # Only header + total rows (max_row == 2)
    assert ws.max_row == 2


# ── PDF ───────────────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_pdf_export_returns_correct_content_type(client: AsyncClient):
    h = await _auth(client, "pdf1")
    with patch(
        "app.services.export_service.generate_pdf",
        new=AsyncMock(return_value=b"%PDF-fake"),
    ):
        resp = await client.get(
            f"{BASE}/expenses/export", params={"format": "pdf", "month": MONTH}, headers=h
        )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"


@pytest.mark.asyncio
async def test_pdf_export_content_disposition(client: AsyncClient):
    h = await _auth(client, "pdf2")
    with patch(
        "app.services.export_service.generate_pdf",
        new=AsyncMock(return_value=b"%PDF-fake"),
    ):
        resp = await client.get(
            f"{BASE}/expenses/export", params={"format": "pdf", "month": MONTH}, headers=h
        )
    assert MONTH in resp.headers["content-disposition"]
    assert resp.headers["content-disposition"].endswith(".pdf\"")


# ── Invalid format ────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_invalid_format_returns_400(client: AsyncClient):
    h = await _auth(client, "bad1")
    resp = await client.get(
        f"{BASE}/expenses/export", params={"format": "csv", "month": MONTH}, headers=h
    )
    assert resp.status_code == 400


# ── Default month (no param) ──────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_excel_export_default_month(client: AsyncClient):
    h = await _auth(client, "def1")
    resp = await client.get(f"{BASE}/expenses/export", params={"format": "xlsx"}, headers=h)
    assert resp.status_code == 200
    assert "attachment" in resp.headers["content-disposition"]


# ── Auth guard ────────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_export_requires_auth(client: AsyncClient):
    resp = await client.get(f"{BASE}/expenses/export", params={"format": "xlsx", "month": MONTH})
    assert resp.status_code == 401
