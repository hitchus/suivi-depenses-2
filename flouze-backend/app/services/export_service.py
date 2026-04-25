import asyncio
import io
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import openpyxl
from jinja2 import Environment, FileSystemLoader
from sqlalchemy import and_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.budget import Budget
from app.models.category import Category
from app.models.expense import Expense
from app.models.user import User

_TEMPLATE_DIR = Path(__file__).parent.parent / "templates"


async def _fetch_expense_rows(
    db: AsyncSession,
    user_id,
    month: str,
) -> tuple[list[dict], Decimal]:
    """Return (rows, total_mad) for the given user/month."""
    year_n, month_n = (int(p) for p in month.split("-"))
    from sqlalchemy import extract, func

    exps = (
        await db.execute(
            select(Expense)
            .where(
                and_(
                    Expense.deleted_at.is_(None),
                    Expense.user_id == user_id,
                    extract("year", Expense.date) == year_n,
                    extract("month", Expense.date) == month_n,
                )
            )
            .order_by(Expense.date.desc(), Expense.created_at.desc())
        )
    ).scalars().all()

    # Build category name map
    cat_ids = {e.category_id for e in exps if e.category_id}
    cat_map: dict = {}
    if cat_ids:
        cats = (
            await db.execute(select(Category).where(Category.id.in_(cat_ids)))
        ).scalars().all()
        cat_map = {c.id: c.name for c in cats}

    rows = [
        {
            "date": e.date.isoformat(),
            "title": e.title,
            "category": cat_map.get(e.category_id, "") if e.category_id else "",
            "amount": str(e.amount),
            "currency": e.currency,
            "amount_mad": str(e.amount_mad),
            "note": e.note or "",
        }
        for e in exps
    ]
    total = sum(e.amount_mad for e in exps) or Decimal("0")
    return rows, total


async def generate_excel(
    db: AsyncSession,
    user_id,
    month: str,
) -> bytes:
    rows, total = await _fetch_expense_rows(db, user_id, month)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = month

    headers = ["Date", "Titre", "Catégorie", "Montant", "Devise", "Montant (MAD)", "Note"]
    ws.append(headers)

    # Bold header row
    from openpyxl.styles import Font, PatternFill, Alignment
    header_fill = PatternFill("solid", fgColor="2D5F3F")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append([
            row["date"],
            row["title"],
            row["category"],
            row["amount"],
            row["currency"],
            row["amount_mad"],
            row["note"],
        ])

    # Total row
    ws.append(["", "TOTAL", "", "", "", str(total), ""])
    total_row = ws.max_row
    for cell in ws[total_row]:
        cell.font = Font(bold=True)

    # Auto-fit column widths (approximate)
    col_widths = [12, 35, 20, 12, 8, 14, 30]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def generate_pdf(
    db: AsyncSession,
    user_id,
    month: str,
) -> bytes:
    rows, total = await _fetch_expense_rows(db, user_id, month)

    user = await db.get(User, user_id)
    user_currency = user.currency if user else "MAD"
    user_name = user.display_name if user else ""

    # Global budget for the month
    budget_row = (
        await db.execute(
            select(Budget).where(
                and_(
                    Budget.user_id == user_id,
                    Budget.category_id.is_(None),
                    Budget.month == month,
                )
            )
        )
    ).scalar_one_or_none()

    budget = str(budget_row.amount) if budget_row else None
    restant = (budget_row.amount - total) if budget_row else None

    env = Environment(loader=FileSystemLoader(str(_TEMPLATE_DIR)), autoescape=True)
    tmpl = env.get_template("expense_report.html")
    html = tmpl.render(
        month=month,
        user_name=user_name,
        currency=user_currency,
        total=str(total.quantize(Decimal("0.01"))),
        expenses=rows,
        budget=budget,
        restant=str(restant.quantize(Decimal("0.01"))) if restant is not None else None,
        generated_at=datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC"),
    )

    # WeasyPrint is CPU-bound; run in thread pool to avoid blocking the event loop
    def _render(html_str: str) -> bytes:
        from weasyprint import HTML
        return HTML(string=html_str).write_pdf()

    loop = asyncio.get_event_loop()
    return await loop.run_in_executor(None, _render, html)
